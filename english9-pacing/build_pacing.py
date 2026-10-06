"""Build the English 9 Purple/Gold pacing calendar (2026-27) as an .xlsx
that converts cleanly to Google Sheets.

Usage: python3 build_pacing.py <folder with the unit .docx files> <output.xlsx>
"""
import datetime as dt
import glob
import os
import re
import sys

import docx
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

D = dt.date


def days(a, b):
    while a <= b:
        yield a
        a += dt.timedelta(1)


# ---------------------------------------------------------------- calendar
FIRST, LAST = D(2026, 8, 24), D(2027, 6, 8)
CLOSED = {D(2026, 9, 4), D(2026, 9, 7), D(2026, 10, 16), D(2026, 11, 3),
          D(2026, 11, 25), D(2026, 11, 26), D(2026, 11, 27), D(2027, 1, 4),
          D(2027, 1, 18), D(2027, 1, 25), D(2027, 1, 26), D(2027, 2, 15),
          D(2027, 4, 12), D(2027, 5, 31)}
CLOSED |= set(days(D(2026, 12, 21), D(2027, 1, 1)))
CLOSED |= set(days(D(2027, 4, 5), D(2027, 4, 9)))
EXAM = {D(2027, 1, 20), D(2027, 1, 21), D(2027, 1, 22),
        D(2027, 6, 4), D(2027, 6, 7), D(2027, 6, 8)}
HALF = {D(2027, 3, 5): "Half day for students (family conferences)",
        D(2027, 4, 2): "Half day for students; last day of Q3",
        D(2027, 5, 28): "Half day for students"}
QUARTERS = [(D(2026, 11, 2), "Q1"), (D(2027, 1, 22), "Q2"),
            (D(2027, 4, 2), "Q3"), (D(2027, 6, 8), "Q4")]
DAY_NOTES = {D(2026, 8, 24): "First day of school",
             D(2026, 11, 2): "Last day of Q1",
             D(2027, 1, 19): "Last class day before semester exams",
             D(2027, 1, 27): "Second semester begins",
             D(2027, 4, 13): "First day of Q4",
             D(2027, 6, 3): "Last regular class day"}

daytype = {}
toggle = "P"
for d in days(FIRST, LAST):
    if d.weekday() > 4 or d in CLOSED:
        continue
    if d in EXAM:
        daytype[d] = "E"
        continue
    daytype[d] = toggle
    toggle = "G" if toggle == "P" else "P"


def quarter(d):
    for end, q in QUARTERS:
        if d <= end:
            return q
    return "Q4"


# ------------------------------------------------------------- lesson bank
UNIT_TITLES = {1: "Unleashing Influence", 2: "Challenge Accepted!",
               3: "The Drive to Deceive", 4: "Through Thick and Thin",
               5: "The Road Less Traveled", 6: "The Power of Pen & Page",
               7: "Perspectives & Principles"}
UNIT_WEEKS = {1: "5", 2: "4", 3: "5", 4: "4", 5: "9", 6: "4.5", 7: "4.5"}


def clean(s):
    s = s.replace(" ", " ")
    return re.sub(r"\s+", " ", s).strip()


def focus_of(target):
    t = clean(target)
    first = re.split(r"(?<=[.!?])\s+", t)[0]
    first = re.sub(r"^((Today,?\s*)?I will\s+(be able to\s+)?|By the end of class( today)?,?\s*"
                   r"I will be able to\s+)", "", first, flags=re.I)
    return first[:1].upper() + first[1:]


bank = []  # dicts
for path in sorted(glob.glob(os.path.join(sys.argv[1], "*.docx")),
                   key=lambda p: re.search(r"Unit_(\d)", p).group(1)):
    m = re.search(r"Unit_(\d)", path)
    unit = int(m.group(1))
    table = docx.Document(path).tables[1]
    for row in table.rows[1:]:
        c = [cell.text for cell in row.cells]
        label = clean(c[0])
        lm = re.search(r"(Week|Cycle)\s*(\d+).*?Lesson\s*(\d)", label, re.I)
        kind = "W" if lm.group(1).lower() == "week" else "C"
        code = f"{kind}{lm.group(2)}L{lm.group(3)}"
        stds = label[lm.end():]
        stds = re.sub(r"\(Standard|\bWeek\b|\bSOL\b|^[\s/]+|[\s/]+$", "", stds).strip(" /")
        stds = re.sub(r"\s*/\s*", ", ", stds)
        bank.append(dict(
            id=f"U{unit}-{code}", unit=unit, lesson=f"{'Week' if kind == 'W' else 'Cycle'} "
            f"{lm.group(2)}, Lesson {lm.group(3)}", short=f"U{unit} {code}",
            focus=focus_of(c[1]), target=clean(c[1]), stds=stds, eq=clean(c[2]),
            res=clean(c[3]), assess=clean(c[4])[:600], note=""))

by_id = {b["id"]: b for b in bank}

# The Unit 3 mid-quarter now gets its own full day.
b = by_id["U3-C2L5"]
b["focus"] = "Learn poetic devices, then read and annotate “Siren Song” by Margaret Atwood"
b["note"] = "The mid-quarter assessment in this lesson now has its own full day."
by_id["U5-C5L1"]["note"] = ("This lesson is the Q3 benchmark. It is scheduled as the Q3 "
                            "Benchmark assessment day, so it is not in the pacing list.")


def merged(new_id, ids, short, note):
    parts = [by_id[i] for i in ids]
    first = parts[0]
    m = dict(first)
    m.update(id=new_id, short=short,
             lesson=" + ".join(p["lesson"] for p in parts),
             focus=" + ".join(dict.fromkeys(p["focus"] for p in parts)),
             target=" | ".join(dict.fromkeys(p["target"] for p in parts)),
             stds=", ".join(dict.fromkeys(s for p in parts for s in p["stds"].split(", ") if s)),
             note=note)
    bank.append(m)
    by_id[new_id] = m


merged("U1-W1L1+2", ["U1-W1L1", "U1-W1L2"], "U1 W1L1+2",
       "Combined to make room for full-day assessments.")
merged("U2-C2L3+4", ["U2-C2L3", "U2-C2L4"], "U2 C2L3+4",
       "Combined: both lessons have the same learning target.")
merged("U7-C1L2+3", ["U7-C1L2", "U7-C1L3"], "U7 C1L2+3",
       "Combined: finish reading “The Most Dangerous Game” in two blocks instead of three.")

ASSESS = [  # id, short, title, window start, window end, anchor lesson
    ("Q1-MQ", "Q1 MQ", "Q1 English 9 Mid-Quarter (MQ)", D(2026, 9, 23), D(2026, 10, 6), "U1-W4L2"),
    ("Q1-BM", "Q1 Benchmark", "Q1 English 9 Benchmark", D(2026, 10, 28), D(2026, 11, 11), "U2-C3L1"),
    ("Q2-MQ", "Q2 MQ", "Q2 English 9 Mid-Quarter (MQ)", D(2026, 12, 7), D(2026, 12, 18), "U3-C2L5"),
    ("Q2-BM", "Q2 Benchmark", "Q2 English 9 Benchmark", D(2027, 1, 19), D(2027, 2, 3), "U4-C2L1"),
    ("LPA", "LPA", "LPA (Grades 9-11)", D(2027, 2, 8), D(2027, 2, 12), None),
    ("Q3-MQ", "Q3 MQ", "Q3 English 9 Mid-Quarter (MQ)", D(2027, 2, 24), D(2027, 3, 9), "U5-C4L1"),
    ("Q3-BM", "Q3 Benchmark", "Q3 English 9 Benchmark", D(2027, 3, 29), D(2027, 4, 19), "U5-C5L2"),
    ("Q4-MQ", "Q4 MQ", "Q4 English 9 Mid-Quarter (MQ)", D(2027, 5, 17), D(2027, 5, 28), "U6-C2L4"),
    ("Q4-BM", "Q4 Benchmark", "Q4 English 9 Benchmark", D(2027, 6, 3), D(2027, 6, 16), "U7-C2L3"),
]
for aid, short, title, ws, we, _ in ASSESS:
    bank.append(dict(id=aid, unit="", lesson="Assessment (full block)", short=short,
                     focus=title, target="", stds="", eq="", res="",
                     assess=f"Testing window {ws:%m/%d/%Y} to {we:%m/%d/%Y}", note=""))
bank.append(dict(id="FLEX", unit="", lesson="Flex day", short="Flex",
                 focus="Catch-up, reteach, or review", target="", stds="", eq="",
                 res="", assess="", note="Use for catch-up, a snow day, or extra review."))
by_id = {b["id"]: b for b in bank}

SEQUENCE = [b["id"] for b in bank if b["unit"]
            and b["id"] not in {"U1-W1L1", "U1-W1L2", "U2-C2L3", "U2-C2L4",
                                "U7-C1L2", "U7-C1L3", "U5-C5L1"}
            and "+" not in b["id"]]
# put merged lessons where their first part was
order = [b["id"] for b in bank if b["unit"] and "+" not in b["id"]]
pos = {i: n for n, i in enumerate(order)}
pos.update({"U1-W1L1+2": pos["U1-W1L1"], "U2-C2L3+4": pos["U2-C2L3"],
            "U7-C1L2+3": pos["U7-C1L2"]})
SEQUENCE = sorted(SEQUENCE + ["U1-W1L1+2", "U2-C2L3+4", "U7-C1L2+3"], key=pos.get)


# ---------------------------------------------------------------- schedule
def schedule(track, fixed, seq):
    out = dict(fixed)
    it = iter(seq)
    for d in sorted(daytype):
        if daytype[d] != track or d in fixed:
            continue
        out[d] = next(it, "FLEX")
    return out


plans = {}
for track in "PG":
    fixed = {}
    for aid, _, _, ws, we, anchor in ASSESS:
        win = [d for d in sorted(daytype) if daytype[d] == track and ws <= d <= we
               and d not in fixed]
        if not win:  # no class meeting in the window: use the first exam day in it
            win = [d for d in sorted(daytype) if daytype[d] == "E" and ws <= d <= we]
            fixed[win[0]] = aid
            continue
        if anchor is None:
            fixed[win[0]] = aid
            continue
        cur = schedule(track, fixed, SEQUENCE)
        at = next(d for d, v in cur.items() if v == anchor)
        fixed[min(max(at, win[0]), win[-1])] = aid
    plans[track] = schedule(track, fixed, SEQUENCE)
    used = [v for v in plans[track].values() if v in by_id and by_id[v]["unit"]]
    assert used == SEQUENCE, f"{track}: lessons out of order or missing"

# ---------------------------------------------------------------- workbook
FONT = "Arial"
PURPLE, GOLD, EXAMC, MOD, NOSCH = "5B2A86", "F6C244", "52F2DA", "BDF59A", "F47F6B"
HDR_FILL = PatternFill("solid", fgColor="4A3F35")
INPUT_FILL = PatternFill("solid", fgColor="FFF6C8")
thin = Side(style="thin", color="B8AFA3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def fill(c):
    return PatternFill("solid", fgColor=c)


def header(ws, row, labels, widths):
    for i, (lab, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row, i, lab)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row + 1, 1)


wb = Workbook()

# --- Read Me
rm = wb.active
rm.title = "Read Me"
rm.column_dimensions["A"].width = 110
lines = [
    ("English 9 Pacing Calendar, 2026-27 (Purple and Gold days)", True),
    ("Menchville High School. Lessons come from the English 9 Unit 1-7 instructional views. "
     "Day types come from the MHS Purple-Gold calendar; assessment windows come from the NNPS "
     "High School Testing Calendar.", False),
    ("", False),
    ("Tabs", True),
    ("Year at a Glance: month grids colored like the Purple-Gold calendar. Each school day shows "
     "the lesson or assessment for the classes that meet that day. It fills itself from the "
     "Purple and Gold tabs, so do not type in it.", False),
    ("Purple / Gold: one row per class meeting. The Lesson ID column (yellow) is the only column "
     "you need to edit. Unit, lesson, focus and standards fill in from the Lesson Bank.", False),
    ("Lesson Bank: every lesson from the unit documents, plus the assessments and a Flex day. "
     "Add a row here to create a new lesson ID.", False),
    ("Assessments: where each midquarter, benchmark and LPA landed for each track, inside its "
     "testing window.", False),
    ("", False),
    ("How to edit", True),
    ("Change one day: click its Lesson ID cell and pick another ID from the dropdown.", False),
    ("Push lessons back a day (snow day, assembly): on the Purple or Gold tab, select the Lesson "
     "ID cells from that day to the end of the list, cut (Ctrl+X), click the cell one row down "
     "and paste (Ctrl+V). Type FLEX or a new ID in the empty cell. Move any assessment rows "
     "back to their dates afterward.", False),
    ("Add your own notes in the My notes column; it is not linked to anything.", False),
    ("", False),
    ("Colors", True),
    ("Purple = Purple (odd) day, Gold = Gold (even) day, Teal = Exam day, Light green = half day, "
     "Salmon = no school. On the Purple and Gold tabs, assessment rows are teal.", False),
    ("", False),
    ("Scheduling decisions", True),
    ("Each track has 86 class days. Nine full-day assessments plus 80 lessons did not fit, so "
     "three pairs of lessons are combined (see the Lesson Bank notes): U1 W1L1+2, U2 C2L3+4 "
     "and U7 C1L2+3. Unit 5 Cycle 5 Lesson 1 is the Q3 benchmark itself.", False),
    ("Purple classes do not meet during the Q4 benchmark window (6/3-6/16) before exams, so the "
     "Purple Q4 benchmark is on the 6/4 exam day. Confirm this with your department.", False),
    ("Purple ends with one Flex day because Purple has one fewer assessment on a regular day.", False),
]
for r, (t, bold) in enumerate(lines, 1):
    c = rm.cell(r, 1, t)
    c.font = Font(name=FONT, bold=bold, size=14 if r == 1 else 11)
    c.alignment = Alignment(wrap_text=True, vertical="top")

# --- Lesson Bank
lb = wb.create_sheet("Lesson Bank")
cols = ["Lesson ID", "Unit", "Unit title", "Lesson", "Calendar label", "Lesson focus",
        "Learning target", "Standards", "Essential question", "Core resources",
        "Assessment notes", "Scheduling note"]
header(lb, 1, cols, [13, 6, 22, 22, 12, 45, 60, 26, 30, 22, 40, 34])
for r, b in enumerate(bank, 2):
    vals = [b["id"], b["unit"], UNIT_TITLES.get(b["unit"], ""), b["lesson"], b["short"],
            b["focus"], b["target"], b["stds"], b["eq"], b["res"], b["assess"], b["note"]]
    for ci, v in enumerate(vals, 1):
        c = lb.cell(r, ci, v)
        c.font = Font(name=FONT, size=10)
        c.alignment = WRAP
        c.border = BORDER
    if not b["unit"]:
        for ci in range(1, len(cols) + 1):
            lb.cell(r, ci).fill = fill("D9FBF5")
lb.auto_filter.ref = f"A1:L{len(bank) + 1}"

# --- Purple / Gold tabs
TRACK_NAME = {"P": "Purple", "G": "Gold"}
for track in "PG":
    name = TRACK_NAME[track]
    ws = wb.create_sheet(name)
    cols = ["Date", "Day", "Quarter", "Class #", "Lesson ID", "Unit", "Lesson",
            "Lesson focus", "Standards", "Calendar label", "Calendar notes", "My notes"]
    header(ws, 1, cols, [11, 6, 8, 7, 13, 22, 22, 50, 24, 13, 30, 30])
    ws.cell(1, 5).fill = fill("B8860B")
    rows = sorted(plans[track].items())
    for r, (d, lid) in enumerate(rows, 2):
        note = DAY_NOTES.get(d, "")
        if d in HALF:
            note = HALF[d]
        if daytype[d] == "E":
            note = "Exam day (half day). Benchmark window has no other Purple class."
        vals = [d, d.strftime("%a"), quarter(d), "=ROW()-1", lid,
                f'=IFERROR(VLOOKUP($E{r},\'Lesson Bank\'!$A:$L,3,FALSE)&"","")',
                f'=IFERROR(VLOOKUP($E{r},\'Lesson Bank\'!$A:$L,4,FALSE)&"","")',
                f'=IFERROR(VLOOKUP($E{r},\'Lesson Bank\'!$A:$L,6,FALSE)&"","")',
                f'=IFERROR(VLOOKUP($E{r},\'Lesson Bank\'!$A:$L,8,FALSE)&"","")',
                f'=IFERROR(VLOOKUP($E{r},\'Lesson Bank\'!$A:$L,5,FALSE),$E{r})',
                note, ""]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(r, ci, v)
            c.font = Font(name=FONT, size=10)
            c.alignment = WRAP
            c.border = BORDER
        ws.cell(r, 1).number_format = "m/d/yyyy"
        ws.cell(r, 5).fill = INPUT_FILL
        ws.cell(r, 1).fill = fill(PURPLE if track == "P" else GOLD)
        ws.cell(r, 1).font = Font(name=FONT, size=10, bold=True,
                                  color="FFFFFF" if track == "P" else "000000")
        if d in HALF:
            ws.cell(r, 11).fill = fill(MOD)
        if daytype[d] == "E":
            ws.cell(r, 1).fill = fill(EXAMC)
            ws.cell(r, 1).font = Font(name=FONT, size=10, bold=True)
    last = len(rows) + 1
    ws.conditional_formatting.add(
        f"B2:D{last + 200}",
        FormulaRule(formula=['OR(LEFT($E2,1)="Q",$E2="LPA")'], fill=fill(EXAMC)))
    ws.conditional_formatting.add(
        f"F2:J{last + 200}",
        FormulaRule(formula=['OR(LEFT($E2,1)="Q",$E2="LPA")'], fill=fill(EXAMC),
                    font=Font(bold=True)))
    dv = DataValidation(type="list", formula1="='Lesson Bank'!$A$2:$A$300", allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"E2:E{last + 200}")

# --- Year at a Glance
ya = wb.create_sheet("Year at a Glance", 1)
ya.sheet_view.showGridLines = False
months = [(2026, 8), (2026, 9), (2026, 10), (2026, 11), (2026, 12), (2027, 1),
          (2027, 2), (2027, 3), (2027, 4), (2027, 5), (2027, 6)]
c = ya.cell(1, 1, "English 9 Purple-Gold Pacing Calendar, August 2026 to June 2027")
c.font = Font(name=FONT, bold=True, size=14)
legend = [("Purple day", PURPLE, "FFFFFF"), ("Gold day", GOLD, "000000"),
          ("Exam day", EXAMC, "000000"), ("Half day", MOD, "000000"),
          ("No school", NOSCH, "000000")]
for i, (lab, bg, fg) in enumerate(legend):
    cc = ya.cell(2, 1 + i * 3, lab)
    ya.merge_cells(start_row=2, start_column=1 + i * 3, end_row=2, end_column=3 + i * 3)
    cc.fill = fill(bg)
    cc.font = Font(name=FONT, bold=True, size=9, color=fg)
    cc.alignment = CENTER
ya.cell(3, 1, "Each cell shows the date and the lesson for the classes that meet that day. "
        "Edit lessons on the Purple and Gold tabs.").font = Font(name=FONT, italic=True, size=9)

PER_ROW = 3
BLOCK_W = 8  # 7 days + spacer
for col in range(1, PER_ROW * BLOCK_W + 1):
    ya.column_dimensions[get_column_letter(col)].width = 2 if col % BLOCK_W == 0 else 10.5
top = 5
import calendar as cal
for mi, (y, m) in enumerate(months):
    band, slot = divmod(mi, PER_ROW)
    r0 = top + band * 9
    c0 = 1 + slot * BLOCK_W
    t = ya.cell(r0, c0, D(y, m, 1).strftime("%B %Y"))
    ya.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 6)
    t.fill = HDR_FILL
    t.font = Font(name=FONT, bold=True, color="FFFFFF")
    t.alignment = CENTER
    for i, dn in enumerate("SMTWTFS"):
        h = ya.cell(r0 + 1, c0 + i, dn)
        h.font = Font(name=FONT, bold=True, size=9)
        h.alignment = CENTER
        h.border = BORDER
    first = (cal.weekday(y, m, 1) + 1) % 7
    for day in range(1, cal.monthrange(y, m)[1] + 1):
        d = D(y, m, day)
        wk, wd = divmod(first + day - 1, 7)
        cell = ya.cell(r0 + 2 + wk, c0 + wd)
        cell.border = BORDER
        cell.alignment = CENTER
        dt_type = daytype.get(d)
        dexpr = f"DATE({y},{m},{day})"
        if dt_type in ("P", "G"):
            tab = TRACK_NAME[dt_type]
            cell.value = f'="{day}"&CHAR(10)&IFERROR(VLOOKUP({dexpr},{tab}!$A:$J,10,FALSE),"")'
            cell.fill = fill(MOD if d in HALF else (PURPLE if dt_type == "P" else GOLD))
            white = dt_type == "P" and d not in HALF
            cell.font = Font(name=FONT, size=8, bold=True, color="FFFFFF" if white else "000000")
        elif dt_type == "E":
            cell.value = (f'="{day}"&IFERROR(CHAR(10)&"P: "&VLOOKUP({dexpr},Purple!$A:$J,10,FALSE),"")'
                          f'&IFERROR(CHAR(10)&"G: "&VLOOKUP({dexpr},Gold!$A:$J,10,FALSE),"")'
                          f'&IF(COUNTIF(Purple!$A:$A,{dexpr})+COUNTIF(Gold!$A:$A,{dexpr})=0,'
                          f'CHAR(10)&"Exam day","")')
            cell.fill = fill(EXAMC)
            cell.font = Font(name=FONT, size=8, bold=True)
        else:
            cell.value = str(day)
            cell.fill = fill(NOSCH)
            cell.font = Font(name=FONT, size=8, color="7A1F12")
    for wk in range(6):
        ya.row_dimensions[r0 + 2 + wk].height = 28
        for wd in range(7):
            ya.cell(r0 + 2 + wk, c0 + wd).border = BORDER

# --- Assessments
asx = wb.create_sheet("Assessments")
header(asx, 1, ["Assessment", "Window start", "Window end", "Purple date", "Gold date",
                "Planned near (unit guide)"], [32, 13, 13, 13, 13, 40])
for r, (aid, short, title, ws_, we, anchor) in enumerate(ASSESS, 2):
    pd_ = next(d for d, v in plans["P"].items() if v == aid)
    gd_ = next(d for d, v in plans["G"].items() if v == aid)
    near = "" if anchor is None else f"{by_id[anchor]['short']}: {by_id[anchor]['focus'][:60]}"
    for ci, v in enumerate([title, ws_, we, pd_, gd_, near], 1):
        c = asx.cell(r, ci, v)
        c.font = Font(name=FONT, size=10)
        c.border = BORDER
        c.alignment = WRAP
        if isinstance(v, dt.date):
            c.number_format = "m/d/yyyy"
    asx.cell(r, 4).fill = fill(EXAMC if daytype[pd_] == "E" else "E6DDF2")
    asx.cell(r, 5).fill = fill("FCEFC7")
r = len(ASSESS) + 3
asx.cell(r, 1, "Other English testing (from the NNPS testing calendar)").font = \
    Font(name=FONT, bold=True)
other = [("LPA (Grade 9 retake, students who need it)", "11/16/2026", "11/20/2026")]
for i, (n, a, b2) in enumerate(other, r + 1):
    for ci, v in enumerate([n, a, b2], 1):
        asx.cell(i, ci, v).font = Font(name=FONT, size=10)


wb.save(sys.argv[2])

# console summary
for track in "PG":
    print(TRACK_NAME[track], {a[0]: str(next(d for d, v in plans[track].items() if v == a[0]))
                              for a in ASSESS})
    starts = {}
    for d, v in sorted(plans[track].items()):
        u = by_id[v]["unit"]
        if u and u not in starts:
            starts[u] = str(d)
    print("  unit starts", starts, "flex", sum(v == "FLEX" for v in plans[track].values()))
